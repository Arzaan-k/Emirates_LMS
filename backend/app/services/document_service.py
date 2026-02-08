"""
Document Text Extraction Service
Extracts text from PDF, Word (DOCX), and PowerPoint (PPTX) files
"""

import logging
import os
from typing import Optional

logger = logging.getLogger(__name__)


class DocumentService:
    """Service for extracting text from various document formats."""

    @staticmethod
    def extract_text_from_pdf(file_path: str) -> Optional[str]:
        """
        Extract text from PDF file using PyPDF2.

        Args:
            file_path: Path to PDF file

        Returns:
            Extracted text or None if extraction fails
        """
        try:
            import PyPDF2

            text_content = []

            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                num_pages = len(pdf_reader.pages)

                logger.info(f"Extracting text from PDF: {num_pages} pages")

                for page_num in range(num_pages):
                    page = pdf_reader.pages[page_num]
                    text = page.extract_text()
                    if text:
                        text_content.append(text)

            full_text = "\n\n".join(text_content)
            logger.info(f"PDF extraction complete: {len(full_text)} characters")
            return full_text if full_text.strip() else None

        except Exception as e:
            logger.error(f"Error extracting text from PDF: {e}")
            return None

    @staticmethod
    def extract_text_from_docx(file_path: str) -> Optional[str]:
        """
        Extract text from Word DOCX file.

        Args:
            file_path: Path to DOCX file

        Returns:
            Extracted text or None if extraction fails
        """
        try:
            from docx import Document

            doc = Document(file_path)
            text_content = []

            # Extract text from paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content.append(paragraph.text)

            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text_content.append(cell.text)

            full_text = "\n\n".join(text_content)
            logger.info(f"DOCX extraction complete: {len(full_text)} characters")
            return full_text if full_text.strip() else None

        except Exception as e:
            logger.error(f"Error extracting text from DOCX: {e}")
            return None

    @staticmethod
    def extract_text_from_pptx(file_path: str) -> Optional[str]:
        """
        Extract text from PowerPoint PPTX file.

        Args:
            file_path: Path to PPTX file

        Returns:
            Extracted text or None if extraction fails
        """
        try:
            from pptx import Presentation

            prs = Presentation(file_path)
            text_content = []

            logger.info(f"Extracting text from PPTX: {len(prs.slides)} slides")

            for slide_num, slide in enumerate(prs.slides, start=1):
                slide_text = []

                # Extract text from all shapes in the slide
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text)

                if slide_text:
                    text_content.append(f"[Slide {slide_num}]\n" + "\n".join(slide_text))

            full_text = "\n\n".join(text_content)
            logger.info(f"PPTX extraction complete: {len(full_text)} characters")
            return full_text if full_text.strip() else None

        except Exception as e:
            logger.error(f"Error extracting text from PPTX: {e}")
            return None

    @staticmethod
    def extract_text_from_file(file_path: str, file_extension: str) -> Optional[str]:
        """
        Extract text from document based on file extension.

        Args:
            file_path: Path to document file
            file_extension: File extension (e.g., '.pdf', '.docx', '.pptx')

        Returns:
            Extracted text or None if extraction fails or unsupported format
        """
        if not os.path.exists(file_path):
            logger.error(f"File not found: {file_path}")
            return None

        ext = file_extension.lower()

        try:
            if ext == '.pdf':
                return DocumentService.extract_text_from_pdf(file_path)
            elif ext in ['.docx', '.doc']:
                # Note: .doc (old format) requires additional conversion
                # For now, we only fully support .docx
                if ext == '.doc':
                    logger.warning("DOC format detected. Only DOCX is fully supported.")
                    return None
                return DocumentService.extract_text_from_docx(file_path)
            elif ext in ['.pptx', '.ppt']:
                # Note: .ppt (old format) requires additional conversion
                # For now, we only fully support .pptx
                if ext == '.ppt':
                    logger.warning("PPT format detected. Only PPTX is fully supported.")
                    return None
                return DocumentService.extract_text_from_pptx(file_path)
            else:
                logger.warning(f"Unsupported document format: {ext}")
                return None

        except Exception as e:
            logger.error(f"Error extracting text from {ext} file: {e}")
            return None

    @staticmethod
    def should_process_document(resource_type: str) -> bool:
        """
        Check if document type should be processed for text extraction.

        Args:
            resource_type: The resource type (e.g., 'Document', 'Presentation')

        Returns:
            True if document should be processed, False otherwise
        """
        processable_types = ['Document', 'Presentation', 'PDF']
        return resource_type in processable_types
